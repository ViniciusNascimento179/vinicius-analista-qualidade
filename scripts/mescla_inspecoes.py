#!/usr/bin/env python3
"""Mescla uma planilha AVULSA de Inspecoes das Instalacoes na principal.

Se existir em data/ um .xlsx cujo nome contenha 'inspec' (e que nao seja a
planilha principal), o conteudo da aba de inspecao e copiado para a planilha
principal, substituindo a aba correspondente. Depois a avulsa e apagada.

Valida os cabecalhos antes de sobrescrever: se faltar coluna, aborta com
mensagem clara em vez de gerar um painel sem dados.
"""
from __future__ import annotations

import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
PRINCIPAL = DATA_DIR / "planos_de_acao_qualidade.xlsx"
ABA = "Inspeção das Instalações"


def _sem_acento(txt) -> str:
    return unicodedata.normalize("NFKD", str(txt)).encode("ascii", "ignore").decode()


def _chave(txt) -> str:
    return _sem_acento(txt).strip().strip("'\"").lower().replace(" ", "")


def achar_avulsa() -> Path | None:
    for p in sorted(DATA_DIR.glob("*")):
        if (
            p.is_file()
            and p.suffix.lower() in (".xlsx", ".xlsm")
            and "inspec" in _sem_acento(p.stem).lower()
            and p.name.lower() != PRINCIPAL.name.lower()
            and not p.name.startswith("~$")
        ):
            return p
    return None


def linha_cabecalho(ws) -> int:
    for i, linha in enumerate(ws.iter_rows(values_only=True), start=1):
        if sum(1 for v in linha if v is not None) >= 3:
            return i
    return 1


def main() -> None:
    avulsa = achar_avulsa()
    if not avulsa:
        print("[mescla] nenhuma planilha avulsa de inspecoes em data/ - nada a fazer")
        return

    if not PRINCIPAL.exists():
        raise SystemExit(f"[mescla] ERRO: planilha principal nao encontrada ({PRINCIPAL.name})")

    src_wb = openpyxl.load_workbook(avulsa, data_only=True)
    ws_src = None
    for nome in src_wb.sheetnames:
        if "inspec" in _sem_acento(nome).lower():
            ws_src = src_wb[nome]
            break
    if ws_src is None:
        ws_src = src_wb[src_wb.sheetnames[0]]

    linhas = list(ws_src.iter_rows(values_only=True))
    ini = linha_cabecalho(ws_src) - 1
    cab_src = [_chave(c) for c in linhas[ini] if c is not None]
    corpo = linhas[ini + 1:]
    print(f"[mescla] origem: {avulsa.name} / aba '{ws_src.title}' "
          f"({len(corpo)} linhas apos o cabecalho)")

    dst_wb = openpyxl.load_workbook(PRINCIPAL)
    alvo = None
    for nome in dst_wb.sheetnames:
        if _chave(nome) == _chave(ABA):
            alvo = dst_wb[nome]
            break
    if alvo is None:
        raise SystemExit(f"[mescla] ERRO: aba '{ABA}' nao existe em {PRINCIPAL.name}")

    ini_dst = linha_cabecalho(alvo) - 1
    cab_dst = [_chave(c) for c in next(alvo.iter_rows(min_row=ini_dst + 1,
                                                      max_row=ini_dst + 1,
                                                      values_only=True)) if c is not None]
    faltando = [c for c in cab_dst if c not in cab_src]
    if faltando:
        raise SystemExit(
            "[mescla] ERRO: a planilha avulsa nao tem as colunas esperadas.\n"
            f"         Faltando: {faltando}\n"
            "         Use o modelo com as mesmas colunas da aba de Inspecoes.")

    alvo.delete_rows(ini_dst + 2, max(alvo.max_row - ini_dst - 1, 0))
    escritas = 0
    for linha in corpo:
        if all(v is None for v in linha):
            continue
        valores = list(linha)[:len(cab_dst)]
        valores += [None] * (len(cab_dst) - len(valores))
        alvo.append(valores)
        escritas += 1

    dst_wb.save(PRINCIPAL)
    avulsa.unlink()
    print(f"[mescla] OK: {escritas} linhas gravadas na aba '{alvo.title}' de {PRINCIPAL.name}")
    print(f"[mescla] planilha avulsa removida: {avulsa.name}")


if __name__ == "__main__":
    main()
